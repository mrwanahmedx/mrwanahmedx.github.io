"""Reproducible checks on synthetic data; no external services."""
import csv
import json
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import reconcile

from openpyxl import load_workbook


class ReconciliationTests(unittest.TestCase):
    def setUp(self):
        self.orders = ROOT / "sample_input" / "orders.csv"
        self.payments = ROOT / "sample_input" / "payments.csv"

    def test_demo_counts(self):
        rows, rejected, counts = reconcile.reconcile(self.orders, self.payments)
        self.assertEqual(len(rows), 8)
        self.assertEqual(len(rejected), 2)
        self.assertEqual(counts, {
            "matched": 2, "amount_mismatch": 1, "missing_payment": 1,
            "orphan_payment": 1, "duplicate_order": 1,
            "duplicate_payment": 1, "pending_excluded": 1,
            "unexpected_pending_payment": 0,
        })

    def test_duplicate_rows_never_count_as_matches(self):
        rows, _, _ = reconcile.reconcile(self.orders, self.payments)
        status = {r["order_id"]: r["status"] for r in rows}
        self.assertEqual(status["O-1001"], "duplicate_payment")
        self.assertEqual(status["O-1005"], "duplicate_order")
        self.assertEqual(status["O-1007"], "matched")

    def test_amount_difference_preserves_cents(self):
        rows, _, _ = reconcile.reconcile(self.orders, self.payments)
        bad = next(x for x in rows if x["order_id"] == "O-1002")
        self.assertEqual(bad["difference"], "-1.00")
        self.assertEqual(reconcile.money(reconcile.currency("$1,200.50")), "1200.50")
        for value in ("NaN", "-2", "1.999", ""):
            with self.assertRaises(ValueError):
                reconcile.currency(value)

    def test_rejected_rows_keep_provenance(self):
        _, rejected, _ = reconcile.reconcile(self.orders, self.payments)
        actual = {(r["dataset"], r["source_row"], r["order_id"]) for r in rejected}
        self.assertEqual(actual, {("orders", 8, "O-1006"),
                                  ("payments", 9, "O-1009")})

    def test_required_columns_fail_clearly(self):
        with tempfile.TemporaryDirectory() as temp:
            bad = Path(temp) / "bad.csv"
            bad.write_text("order_id\nO-1001\n", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "missing required columns"):
                reconcile.reconcile(bad, self.payments)

    def test_real_excel_and_csv_outputs(self):
        with tempfile.TemporaryDirectory() as temp:
            out = Path(temp)
            summary = reconcile.run(self.orders, self.payments, out)
            self.assertEqual(summary["counts"]["matched"], 2)
            self.assertEqual(sorted(x.name for x in out.iterdir()),
                             ["exceptions.csv", "matched.csv", "reconciliation.xlsx",
                              "rejected_rows.csv", "summary.json"])
            with (out / "matched.csv").open(newline="", encoding="utf-8") as fh:
                matches = list(csv.DictReader(fh))
            self.assertEqual({r["order_id"] for r in matches}, {"O-1007", "O-1008"})
            wb = load_workbook(out / "reconciliation.xlsx", read_only=True)
            self.assertEqual(wb.sheetnames,
                             ["Overview", "Matches", "Exceptions", "Pending", "Rejected Rows"])
            self.assertEqual(wb["Matches"].max_row, 3)
            self.assertEqual(wb["Exceptions"].max_row, 6)
            self.assertEqual(wb["Rejected Rows"].max_row, 3)
            self.assertEqual(json.loads((out / "summary.json").read_text())[
                "rejected_input_rows"], 2)
            wb.close()

    def test_unexpected_payment_on_pending_order(self):
        with tempfile.TemporaryDirectory() as temp:
            csv_path = Path(temp) / "pays.csv"
            csv_path.write_text("order_id,received_amount,reference\n"
                                "O-1004,30.00,SET-P\n", encoding="utf-8")
            rows, _, counts = reconcile.reconcile(self.orders, csv_path)
            status = {r["order_id"]: r["status"] for r in rows}
            self.assertEqual(status["O-1004"], "unexpected_pending_payment")
            self.assertEqual(counts["unexpected_pending_payment"], 1)

    def test_spreadsheet_formula_text_is_escaped(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            orders = root / "orders.csv"
            payments = root / "payments.csv"
            orders.write_text("order_id,customer,expected_amount,status\n"
                              "O-1,=1+1,10.00,completed\n", encoding="utf-8")
            payments.write_text("order_id,received_amount,reference\n"
                                "O-1,10.00,R1\n", encoding="utf-8")
            out = root / "out"
            reconcile.run(orders, payments, out)
            with (out / "matched.csv").open(newline="", encoding="utf-8") as fh:
                self.assertEqual(next(csv.DictReader(fh))["customer"], "'=1+1")
            wb = load_workbook(out / "reconciliation.xlsx")
            self.assertEqual(wb["Matches"]["B2"].value, "'=1+1")
            self.assertNotEqual(wb["Matches"]["B2"].data_type, "f")
            wb.close()


if __name__ == "__main__":
    unittest.main()
